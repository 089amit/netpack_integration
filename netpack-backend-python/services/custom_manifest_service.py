import os
import io
import re
import uuid
import zipfile
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from num2words import num2words

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'templates', 'CHAMBER_template.docx')
TEMP_OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'temp_outputs')
os.makedirs(TEMP_OUTPUT_DIR, exist_ok=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

def is_url(text: Any) -> bool:
    if not text or not isinstance(text, str):
        return False
    t = text.strip().lower()
    return t.startswith('http://') or t.startswith('https://')

def simplify_service_for_display(service: str) -> str:
    if is_url(service):
        try:
            from urllib.parse import urlparse
            domain = urlparse(service).netloc.replace('www.', '')
            return f"Track on {domain}"
        except Exception:
            return "Tracking Link"
    return service

def is_console_bag_marking(marking: Any) -> bool:
    if not marking:
        return False
    return bool(re.match(r'^\s*BAG\s*-\s*\S+', str(marking).strip(), re.IGNORECASE))

DEFAULT_MAX_BAG_WEIGHT = 30.0  # kg - matches authentic courier rule

def suggest_bag_markings(shipments_data: List[Dict[str, Any]], max_bag_weight: float = DEFAULT_MAX_BAG_WEIGHT) -> Dict[str, str]:
    """
    Auto-suggest a bag/box marking per HAWB (matching old manifest project logic):
      - Small HAWBs (single box, weight under max_bag_weight) are consolidated
        together into one bag, filling it up to ~max_bag_weight (kg); every
        HAWB in that group shares the SAME single marking number (1, 2, 3...).
      - A HAWB with multiple boxes of its own (NO OF PCS > 1), or whose own
        weight already meets/exceeds max_bag_weight, is NOT merged with
        others. Instead it gets its own sequential marking: a single number
        if it's one box, or a number RANGE if it has several boxes (e.g.
        '1-8' for an 8-box shipment).
      - One running counter is shared across the whole manifest, so bag
        numbers and box-range numbers interleave in a single sequence.
    """
    suggestions = {}
    counter = 1
    group_hawbs: List[str] = []
    group_weight = 0.0

    def flush_group():
        nonlocal counter, group_hawbs, group_weight
        if not group_hawbs:
            return
        label = str(counter)
        for h in group_hawbs:
            suggestions[h] = label
        counter += 1
        group_hawbs = []
        group_weight = 0.0

    # Sort shipments by hawb
    sorted_shipments = sorted(
        shipments_data,
        key=lambda s: str(s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or '')
    )

    for s in sorted_shipments:
        hawb = str(s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or '')
        weight = float(s.get('weight') or s.get('actualWeight') or 0.0)
        pcs = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)

        if pcs > 1 or weight >= max_bag_weight:
            flush_group()
            start = counter
            end = counter + pcs - 1
            suggestions[hawb] = str(start) if pcs == 1 else f"{start}-{end}"
            counter = end + 1
            continue

        if group_hawbs and (group_weight + weight > max_bag_weight):
            flush_group()

        group_hawbs.append(hawb)
        group_weight += weight

    flush_group()
    return suggestions

def build_bag_groups(shipments_data: List[Dict[str, Any]], hawb_to_bag: Dict[str, str]) -> Tuple[Dict[str, List[Dict[str, Any]]], List[Dict[str, Any]]]:
    """Group shipments by their assigned bag marking."""
    bag_groups: Dict[str, List[Dict[str, Any]]] = {}
    unassigned = []
    for s in shipments_data:
        hawb = str(s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or '')
        bag = str(hawb_to_bag.get(hawb) or '').strip()
        if bag:
            bag_groups.setdefault(bag, []).append(s)
        else:
            unassigned.append(s)
    return bag_groups, unassigned

def build_bag_rows(bag_groups: Dict[str, List[Dict[str, Any]]], unassigned: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Build rows for manifest partitioning in bag-based mode:
    - If marking is 'BAG-<n>' (console bag): collapse all HAWBs into 1 physical box with summed weight & USD.
    - Otherwise (e.g. '1', '1-8'): each HAWB keeps its real boxes, tagged with 'Bag: <marking>' in remarks.
    """
    bag_rows = []
    for bag, rows_list in bag_groups.items():
        if is_console_bag_marking(bag):
            total_weight = sum(float(r.get('weight') or r.get('actualWeight') or 0.0) for r in rows_list)
            total_usd = sum(float(r.get('value') or r.get('totalValue') or 100.0) for r in rows_list)
            first_row = rows_list[0]
            descriptions = [str(r.get('goodsDescription') or r.get('description') or '') for r in rows_list if r.get('goodsDescription') or r.get('description')]
            combined_desc = ', '.join(dict.fromkeys(descriptions)) if descriptions else 'Consolidated Goods'

            combined_items = []
            for r in rows_list:
                for it in (r.get('enquiryItems') or []):
                    combined_items.append(it)

            bag_rows.append({
                'id': first_row.get('id'),
                'hawbno': bag,
                'noOfBox': 1,
                'weight': round(total_weight, 2),
                'value': round(total_usd, 2),
                'senderName': first_row.get('senderName', 'Shipper'),
                'receiverName': first_row.get('receiverName', 'Consignee'),
                'destinationCountryName': first_row.get('destinationCountryName', 'UK'),
                'forwardingNumber': first_row.get('forwardingNumber', 'N/A'),
                'forwardingServiceName': first_row.get('forwardingServiceName', 'Standard'),
                'goodsDescription': combined_desc,
                'remarks': f"Console Bag: {bag}",
                'enquiryItems': combined_items,
                'status': 'IN_TRANSIT'
            })
        else:
            for row in rows_list:
                r_copy = dict(row)
                r_copy['remarks'] = f"Bag: {bag}"
                bag_rows.append(r_copy)

    for row in unassigned:
        r_copy = dict(row)
        r_copy['remarks'] = ''
        bag_rows.append(r_copy)

    return bag_rows

def prepare_wtbox_data(shipments_data: List[Dict[str, Any]],
                       bag_markings: Optional[Dict[str, str]] = None,
                       separate_clearance_hawbs: Optional[List[str]] = None,
                       max_bag_weight: float = DEFAULT_MAX_BAG_WEIGHT) -> List[Dict[str, Any]]:
    bag_markings = bag_markings or {}
    separate_set = set(separate_clearance_hawbs or [])
    auto_suggestions = suggest_bag_markings(shipments_data, max_bag_weight=max_bag_weight)

    rows = []
    for s in shipments_data:
        hawb = s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or ''
        sender = s.get('senderName') or 'Shipper'
        receiver = s.get('receiverName') or s.get('reciverName') or 'Consignee'
        details = f"{sender} -> {receiver}"
        pcs = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)
        weight = float(s.get('weight') or s.get('actualWeight') or 0.0)
        country = s.get('destinationCountryName') or s.get('country') or ''
        marking = bag_markings.get(hawb) or auto_suggestions.get(hawb, '')
        is_sep = hawb in separate_set

        rows.append({
            'HAWB': hawb,
            'Details': details,
            'Box': pcs,
            'BagMarking': marking,
            'Weight': round(weight, 2),
            'Country': country,
            'SeparateClearance': is_sep
        })
    return rows

def generate_wtbox_excel(wt_rows: List[Dict[str, Any]], reference: str = "report") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'WT-Box'

    headers = ['HAWB', 'Details', 'Box', 'Bag Marking', 'Weight', 'Country', 'Separate Clearance']
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    total_boxes = 0
    total_weight = 0.0

    for r_idx, r in enumerate(wt_rows, 2):
        b = int(r.get('Box', 0))
        w = float(r.get('Weight', 0.0))
        total_boxes += b
        total_weight += w

        ws.cell(row=r_idx, column=1, value=r.get('HAWB', ''))
        ws.cell(row=r_idx, column=2, value=r.get('Details', ''))
        ws.cell(row=r_idx, column=3, value=b)
        ws.cell(row=r_idx, column=4, value=r.get('BagMarking', ''))
        ws.cell(row=r_idx, column=5, value=w)
        ws.cell(row=r_idx, column=6, value=r.get('Country', ''))
        ws.cell(row=r_idx, column=7, value='Yes' if r.get('SeparateClearance') else '')

        for col in range(1, 8):
            ws.cell(row=r_idx, column=col).border = THIN_BORDER
            ws.cell(row=r_idx, column=col).alignment = Alignment(vertical='top')

    tot_row_idx = len(wt_rows) + 2
    ws.cell(row=tot_row_idx, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=tot_row_idx, column=3, value=total_boxes).font = Font(bold=True)
    ws.cell(row=tot_row_idx, column=5, value=round(total_weight, 2)).font = Font(bold=True)
    for col in range(1, 8):
        ws.cell(row=tot_row_idx, column=col).border = THIN_BORDER

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.freeze_panes = 'A2'
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

def split_manifest(shipments: List[Dict[str, Any]],
                   max_box_per_part: int = 10,
                   max_usd_per_part: float = 950.0,
                   max_weight_per_part: float = 1500.0,
                   separate_clearance_hawbs: Optional[List[str]] = None,
                   is_bag_based: bool = False,
                   bag_markings: Optional[Dict[str, str]] = None) -> List[List[Dict[str, Any]]]:
    import math

    separate_set = set(separate_clearance_hawbs or [])
    active_shipments = [s for s in shipments if (s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#')) not in separate_set]
    if not active_shipments:
        return []

    # If bag-based, group into bag-level entries first
    if is_bag_based and bag_markings:
        bag_groups, unassigned = build_bag_groups(active_shipments, bag_markings)
        base_entries = build_bag_rows(bag_groups, unassigned)
    else:
        base_entries = list(active_shipments)

    # Decompose entries that exceed partition caps
    entries_to_pack: List[Dict[str, Any]] = []
    for s in base_entries:
        boxes = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)
        usd = float(s.get('value') or s.get('totalValue') or 100.0)
        weight = float(s.get('weight') or s.get('actualWeight') or 0.0)
        items = list(s.get('enquiryItems') or [])

        needed_by_box = math.ceil(boxes / max_box_per_part) if max_box_per_part > 0 else 1
        needed_by_usd = math.ceil(usd / max_usd_per_part) if max_usd_per_part > 0 else 1
        needed_by_wt = math.ceil(weight / max_weight_per_part) if max_weight_per_part > 0 else 1
        n_chunks = max(needed_by_box, needed_by_usd, needed_by_wt, 1)

        if n_chunks > 1 and (boxes > 1 or len(items) > 1):
            # Divide this shipment across n_chunks sub-parts
            boxes_left = boxes
            weight_left = weight
            items_left = list(items)

            for chunk_idx in range(n_chunks):
                chunks_remaining = n_chunks - chunk_idx
                cur_b = math.ceil(boxes_left / chunks_remaining)
                cur_w = round(weight * (cur_b / boxes), 2) if boxes > 0 else round(weight / n_chunks, 2)
                cur_w = min(cur_w, weight_left)

                cur_items = []
                if items_left:
                    # Allocate items up to max_usd_per_part
                    target_chunk_usd = max_usd_per_part
                    accumulated_usd = 0.0
                    while items_left:
                        next_it = items_left[0]
                        it_val = float(next_it.get('totalValue') or ((next_it.get('quantity') or 1) * (next_it.get('unitPrice') or 0.0)))
                        if (accumulated_usd + it_val <= target_chunk_usd) or (not cur_items) or (chunks_remaining == 1):
                            cur_items.append(items_left.pop(0))
                            accumulated_usd += it_val
                        else:
                            break
                    cur_val = accumulated_usd
                else:
                    cur_val = round(usd / n_chunks, 2)

                boxes_left -= cur_b
                weight_left -= cur_w

                sub_entry = dict(s)
                sub_entry['noOfBox'] = cur_b
                sub_entry['weight'] = cur_w
                sub_entry['value'] = cur_val
                sub_entry['enquiryItems'] = cur_items
                if cur_items:
                    sub_entry['goodsDescription'] = ", ".join([it.get('description', '') for it in cur_items if it.get('description')])
                entries_to_pack.append(sub_entry)
        else:
            entries_to_pack.append(s)

    # Bin-pack entries into parts
    parts: List[List[Dict[str, Any]]] = []
    current_part: List[Dict[str, Any]] = []
    current_boxes = 0
    current_usd = 0.0
    current_weight = 0.0

    for s in entries_to_pack:
        boxes = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)
        usd = float(s.get('value') or s.get('totalValue') or 100.0)
        weight = float(s.get('weight') or s.get('actualWeight') or 0.0)

        would_exceed_boxes = (current_boxes + boxes > max_box_per_part) and (current_boxes > 0)
        would_exceed_usd = (current_usd + usd > max_usd_per_part) and (current_usd > 0)
        would_exceed_weight = (current_weight + weight > max_weight_per_part) and (current_weight > 0)

        if would_exceed_boxes or would_exceed_usd or would_exceed_weight:
            parts.append(current_part)
            current_part = [s]
            current_boxes = boxes
            current_usd = usd
            current_weight = weight
        else:
            current_part.append(s)
            current_boxes += boxes
            current_usd += usd
            current_weight += weight

    if current_part:
        parts.append(current_part)

    return parts

def extract_sheet_matrix(excel_bytes: bytes) -> Dict[str, List[List[str]]]:
    """Parse generated excel bytes into 2D string matrices per worksheet for on-screen preview."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        res = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            grid = []
            for row in ws.iter_rows(values_only=True):
                row_vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in row_vals):
                    grid.append(row_vals)
            res[sheet_name] = grid
        return res
    except Exception as e:
        return {}

def _add_title_block_sheet1(worksheet, total_boxes, total_weight, invoice_date, max_col, invoice_number,
                            consignee_name="ESHIPPER EXPRESS COURIER", consignee_address="", consignee_country="UK"):
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    worksheet['A1'] = 'DOKO EXPORTS PVT LTD'
    worksheet['A1'].font = Font(bold=True, size=14)
    worksheet['A1'].alignment = Alignment(horizontal='center', vertical='center')

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    worksheet['A2'] = 'INVOICE/PACKING LIST'
    worksheet['A2'].font = Font(bold=True, size=12)
    worksheet['A2'].alignment = Alignment(horizontal='center', vertical='center')

    worksheet['A5'] = f'Invoice No:- {invoice_number}'
    worksheet['A5'].font = Font(bold=True)
    worksheet['G5'] = 'EXIM NO:- 6197818390116NP'
    worksheet['G5'].font = Font(bold=True)

    worksheet['A6'] = 'Consignee Name, Address & Country'
    worksheet['A6'].font = Font(bold=True)
    worksheet['G6'] = f'Invoice Date :- {invoice_date}'
    worksheet['G6'].font = Font(bold=True)

    worksheet['G7'] = 'PAN NO: 619781839'
    worksheet['G7'].font = Font(bold=True)

    worksheet['A8'] = 'To:'
    worksheet['A8'].font = Font(bold=True)
    worksheet['B8'] = consignee_name
    worksheet['G8'] = f'Box- {total_boxes} boxes'
    worksheet['G8'].font = Font(bold=True)

    worksheet['B9'] = ', '.join(p for p in [consignee_address, consignee_country] if p)
    worksheet['G9'] = f'Weight: {total_weight:.2f} kgs'
    worksheet['G9'].font = Font(bold=True)

    worksheet['G10'] = 'Shipment by : Air Freight'
    worksheet['G10'].font = Font(bold=True)
    return 14

def _add_title_block_sheet2(worksheet, total_boxes, total_weight, invoice_date, max_col, invoice_number,
                            consignee_name="ESHIPPER EXPRESS COURIER", consignee_address="", consignee_country="UK"):
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    worksheet['A1'] = 'DOKO EXPORTS PVT LTD'
    worksheet['A1'].font = Font(bold=True, size=14)
    worksheet['A1'].alignment = Alignment(horizontal='center', vertical='center')

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    worksheet['A2'] = 'INVOICE/PACKING LIST'
    worksheet['A2'].font = Font(bold=True, size=12)
    worksheet['A2'].alignment = Alignment(horizontal='center', vertical='center')

    worksheet['A5'] = f'Invoice No:- {invoice_number}'
    worksheet['A5'].font = Font(bold=True)
    worksheet['C5'] = f'Invoice Date :- {invoice_date}'
    worksheet['C5'].font = Font(bold=True)

    worksheet['A6'] = 'Consignee Name, Address & Country'
    worksheet['A6'].font = Font(bold=True)
    worksheet['C6'] = 'EXIM NO:-6197818390116NP'
    worksheet['C6'].font = Font(bold=True)

    worksheet['C7'] = 'PAN NO:619781839'
    worksheet['C7'].font = Font(bold=True)

    worksheet['A8'] = 'To:'
    worksheet['A8'].font = Font(bold=True)
    worksheet['B8'] = consignee_name
    worksheet['C8'] = f'Box- {total_boxes} boxes'
    worksheet['C8'].font = Font(bold=True)

    worksheet['B9'] = ', '.join(p for p in [consignee_address, consignee_country] if p)
    worksheet['C9'] = f'Weight: {total_weight:.2f} kgs'
    worksheet['C9'].font = Font(bold=True)

    worksheet['C10'] = 'Shipment by : Air Freight'
    worksheet['C10'].font = Font(bold=True)
    return 14

def _apply_sheet_borders(worksheet, max_row, max_col):
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            worksheet.cell(row=row, column=col).border = thin_border

def _auto_fit_sheet_columns(worksheet, start_row):
    for column in worksheet.columns:
        max_length = 0
        col_letter = get_column_letter(column[0].column)
        for cell in column[start_row-1:]:
            try:
                if cell.value:
                    if isinstance(cell.value, str) and '\n' in cell.value:
                        lines = cell.value.split('\n')
                        cell_length = max(len(line) for line in lines)
                    else:
                        cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except Exception:
                pass
        adjusted_width = min(max(max_length + 2, 8), 50)
        worksheet.column_dimensions[col_letter].width = adjusted_width
        for cell in column[start_row-1:]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

def generate_custom_manifest_excel(part_shipments: List[Dict[str, Any]],
                                   part_number: int,
                                   invoice_number: str,
                                   invoice_date: str,
                                   consignee_info: Dict[str, Any],
                                   bag_markings: Optional[Dict[str, str]] = None,
                                   narration_text: Optional[str] = None) -> bytes:
    bag_markings = bag_markings or {}
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    c_name = consignee_info.get('name') or 'ESHIPPER EXPRESS COURIER'
    c_addr = consignee_info.get('address') or ''
    c_dest = consignee_info.get('country') or 'UK'

    total_boxes = sum(int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1) for s in part_shipments)
    total_weight = sum(float(s.get('weight') or s.get('actualWeight') or 0.0) for s in part_shipments)
    total_usd = sum(float(s.get('value') or s.get('totalValue') or 100.0) for s in part_shipments)

    # ══════════════════════════════════════════════════════════════════════════
    # ─── SHEET 1: Custom Manifest (Exact format of manifest project) ──────────
    # ══════════════════════════════════════════════════════════════════════════
    ws_manifest = wb.create_sheet('Custom Manifest')
    ws_manifest.page_setup.orientation = 'landscape'
    ws_manifest.page_setup.paperSize = 9  # A4
    ws_manifest.page_setup.fitToPage = True
    ws_manifest.page_setup.fitToWidth = 1
    ws_manifest.page_setup.fitToHeight = 0

    manifest_columns = ['BOX', 'HAWB', 'PCS', 'WEIGHT', 'USD', 'SENDER', 'RECEIVER', 'DESCRIPTION', 'DESTINATION', 'REMARKS']
    max_col_m = len(manifest_columns)
    start_row_m = _add_title_block_sheet1(ws_manifest, total_boxes, total_weight, invoice_date, max_col_m, invoice_number,
                                          c_name, c_addr, c_dest)

    # Header row
    for c_idx, col_name in enumerate(manifest_columns, 1):
        cell = ws_manifest.cell(row=start_row_m, column=c_idx, value=col_name)
        cell.font = Font(bold=True)

    curr_row_m = start_row_m + 1
    box_start = 1
    for s in part_shipments:
        hawb = s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or ''
        pcs = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)
        w = float(s.get('weight') or s.get('actualWeight') or 0.0)
        val = float(s.get('value') or s.get('totalValue') or 100.0)
        sender = s.get('senderName') or 'Shipper'
        receiver = s.get('receiverName') or s.get('reciverName') or 'Consignee'
        desc = s.get('goodsDescription') or s.get('description') or 'Handicrafts & General Goods'
        dest = s.get('destinationCountryName') or s.get('country') or c_dest

        marking = bag_markings.get(hawb, '')
        remarks = f"Console Bag: {marking}" if (marking and 'BAG' in marking.upper()) else (f"Bag: {marking}" if marking else "")

        box_end = box_start + pcs - 1
        box_str = f"{box_start}-{box_end}" if pcs > 1 else str(box_start)
        box_start = box_end + 1

        ws_manifest.cell(row=curr_row_m, column=1, value=box_str)
        ws_manifest.cell(row=curr_row_m, column=2, value=hawb)
        ws_manifest.cell(row=curr_row_m, column=3, value=pcs)
        ws_manifest.cell(row=curr_row_m, column=4, value=w)
        ws_manifest.cell(row=curr_row_m, column=5, value=round(val, 2))
        ws_manifest.cell(row=curr_row_m, column=6, value=sender)
        ws_manifest.cell(row=curr_row_m, column=7, value=receiver)
        ws_manifest.cell(row=curr_row_m, column=8, value=desc)
        ws_manifest.cell(row=curr_row_m, column=9, value=dest)
        ws_manifest.cell(row=curr_row_m, column=10, value=remarks)
        curr_row_m += 1

    # Total row
    ws_manifest.cell(row=curr_row_m, column=1, value='TOTAL').font = Font(bold=True)
    ws_manifest.cell(row=curr_row_m, column=3, value=total_boxes).font = Font(bold=True)
    ws_manifest.cell(row=curr_row_m, column=4, value=round(total_weight, 2)).font = Font(bold=True)
    ws_manifest.cell(row=curr_row_m, column=5, value=round(total_usd, 2)).font = Font(bold=True)

    _apply_sheet_borders(ws_manifest, curr_row_m, max_col_m)
    _auto_fit_sheet_columns(ws_manifest, start_row_m)

    last_row_m = curr_row_m
    if narration_text:
        narration_row = last_row_m + 2
        ws_manifest.merge_cells(start_row=narration_row, start_column=1, end_row=narration_row, end_column=max_col_m)
        cell = ws_manifest.cell(row=narration_row, column=1, value=narration_text)
        cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        cell.font = Font(italic=True, size=10)
        line_count = narration_text.count('\n') + 1
        ws_manifest.row_dimensions[narration_row].height = 15 * line_count

    # ══════════════════════════════════════════════════════════════════════════
    # ─── SHEET 2: Custom Packing List (Exact format of manifest project) ──────
    # ══════════════════════════════════════════════════════════════════════════
    ws_packing = wb.create_sheet('Custom Packing List')
    ws_packing.page_setup.orientation = 'portrait'
    ws_packing.page_setup.paperSize = 9  # A4
    ws_packing.page_setup.fitToPage = True
    ws_packing.page_setup.fitToWidth = 1
    ws_packing.page_setup.fitToHeight = 0

    packing_columns = ['HAWB', 'DESCRIPTION', 'PCS', 'UNIT VALUE', 'TOTAL', 'HSCODE']
    max_col_p = len(packing_columns)
    start_row_p = _add_title_block_sheet2(ws_packing, total_boxes, total_weight, invoice_date, max_col_p, invoice_number,
                                          c_name, c_addr, c_dest)

    # Header row
    for c_idx, col_name in enumerate(packing_columns, 1):
        cell = ws_packing.cell(row=start_row_p, column=c_idx, value=col_name)
        cell.font = Font(bold=True)

    curr_row_p = start_row_p + 1
    sum_pcs = 0
    sum_total = 0.0

    for s_idx, s in enumerate(part_shipments):
        hawb = s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or ''
        items = s.get('enquiryItems') or s.get('items') or []
        marking = bag_markings.get(hawb, '')
        note_text = f"Console Bag: {marking}" if (marking and 'BAG' in marking.upper()) else (f"Bag: {marking}" if marking else "")

        if not items:
            pcs = int(s.get('noOfBox') or s.get('pieces') or 1)
            val = float(s.get('value') or s.get('totalValue') or s.get('finalRate') or 100.0)
            u_val = round(val / pcs if pcs > 0 else val, 2)
            desc = s.get('goodsDescription') or 'General Goods'

            ws_packing.cell(row=curr_row_p, column=1, value=hawb)
            ws_packing.cell(row=curr_row_p, column=2, value=desc)
            ws_packing.cell(row=curr_row_p, column=3, value=pcs)
            ws_packing.cell(row=curr_row_p, column=4, value=u_val)
            ws_packing.cell(row=curr_row_p, column=5, value=round(val, 2))
            ws_packing.cell(row=curr_row_p, column=6, value='')
            sum_pcs += pcs
            sum_total += val
            curr_row_p += 1

            if note_text:
                ws_packing.cell(row=curr_row_p, column=1, value=note_text)
                curr_row_p += 1
        else:
            # REAL ITEMS FROM DATABASE!
            n_items = len(items)
            for i_idx, item in enumerate(items):
                desc = item.get('description') or 'General Goods'
                pcs = int(item.get('quantity') or 1)
                u_val = float(item.get('unitPrice') or 0.0)
                tot = float(item.get('totalValue') or (pcs * u_val) or item.get('value') or 0.0)
                if u_val <= 0 and pcs > 0 and tot > 0:
                    u_val = round(tot / pcs, 2)
                hs = item.get('hsCode') or ''

                if i_idx == 0:
                    hawb_cell = hawb
                elif i_idx == 1 and note_text:
                    hawb_cell = note_text
                else:
                    hawb_cell = ''

                ws_packing.cell(row=curr_row_p, column=1, value=hawb_cell)
                ws_packing.cell(row=curr_row_p, column=2, value=desc)
                ws_packing.cell(row=curr_row_p, column=3, value=pcs)
                ws_packing.cell(row=curr_row_p, column=4, value=round(u_val, 2))
                ws_packing.cell(row=curr_row_p, column=5, value=round(tot, 2))
                ws_packing.cell(row=curr_row_p, column=6, value=hs)

                sum_pcs += pcs
                sum_total += tot
                curr_row_p += 1

            if note_text and n_items < 2:
                ws_packing.cell(row=curr_row_p, column=1, value=note_text)
                curr_row_p += 1

        # Blank separator row between shipments
        if s_idx < len(part_shipments) - 1:
            curr_row_p += 1

    # Total row on Sheet 2
    ws_packing.cell(row=curr_row_p, column=1, value='')
    ws_packing.cell(row=curr_row_p, column=2, value='TOTAL').font = Font(bold=True)
    ws_packing.cell(row=curr_row_p, column=3, value=sum_pcs).font = Font(bold=True)
    ws_packing.cell(row=curr_row_p, column=5, value=round(sum_total, 2)).font = Font(bold=True)

    _apply_sheet_borders(ws_packing, curr_row_p, max_col_p)
    _auto_fit_sheet_columns(ws_packing, start_row_p)

    # Remove default sheet if it was created
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

def replace_placeholders_in_doc(doc: Document, replacements: Dict[str, str]) -> int:
    replace_count = 0
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                for run in p.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, str(val))
                        replace_count += 1
                if key in p.text:
                    p.text = p.text.replace(key, str(val))
                    replace_count += 1

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for key, val in replacements.items():
                    if key in cell.text:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, str(val))
                                    replace_count += 1
                            if key in p.text:
                                p.text = p.text.replace(key, str(val))
                                replace_count += 1
    return replace_count

def generate_chamber_docx(consignee_info: Dict[str, Any],
                          invoice_number: str,
                          invoice_date: str,
                          total_usd: float,
                          total_boxes: int,
                          packing_total_pcs: int) -> Optional[bytes]:
    if not os.path.exists(TEMPLATE_PATH):
        return None

    doc = Document(TEMPLATE_PATH)
    c_name = consignee_info.get('name') or 'ESHIPPER EXPRESS COURIER'
    c_addr = consignee_info.get('address') or ''
    c_dest = consignee_info.get('country') or 'UK'
    consignee_parts = [c_name, c_addr, c_dest]
    consignee_details = ', '.join(p.strip() for p in consignee_parts if p and p.strip())

    try:
        usd_words = num2words(total_usd, lang='en', to='currency', currency='USD').upper()
    except Exception:
        usd_words = f"{total_usd:.2f} USD"

    replacements = {
        '{{MANIFEST_TITLE}}': 'DOKO EXPORTS PVT LTD',
        '{{CONSIGNEE_DETAILS}}': consignee_details,
        '{{DATE}}': invoice_date,
        '{{TOTAL_USD}}': f"{total_usd:.2f}",
        '{{TOTAL_USD_WORDS}}': usd_words,
        '{{TOTAL_BOXES}}': str(total_boxes),
        '{{PACKING_TOTAL_PCS}}': str(packing_total_pcs),
        '{{INVOICE_NUMBER}}': invoice_number
    }

    replace_placeholders_in_doc(doc, replacements)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()

def generate_tracking_report_excel(shipments_data: List[Dict[str, Any]], reference: str = "report") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tracking Report'

    headers = ['Tracking Number', 'HAWB', 'Consignee Name', 'Country', 'PCS', 'Status', 'Service']
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r_idx, s in enumerate(shipments_data, 2):
        fwd_no = str(s.get('forwardingNumber') or s.get('forwarding_number') or 'N/A')
        hawb = s.get('hawbno') or s.get('hawbNumber') or s.get('HAWB#') or ''
        consignee = s.get('receiverName') or s.get('reciverName') or s.get('consignee') or ''
        country = s.get('destinationCountryName') or s.get('country') or ''
        pcs = int(s.get('noOfBox') or s.get('pieces') or len(s.get('boxes') or []) or 1)
        status = s.get('status') or 'IN_TRANSIT'
        service = s.get('forwardingServiceName') or s.get('service') or 'Standard'

        c1 = ws.cell(row=r_idx, column=1)
        c1.value = fwd_no
        c1.number_format = '@'
        c1.alignment = Alignment(vertical='top')

        ws.cell(row=r_idx, column=2, value=hawb).alignment = Alignment(vertical='top')
        ws.cell(row=r_idx, column=3, value=consignee).alignment = Alignment(vertical='top')
        ws.cell(row=r_idx, column=4, value=country).alignment = Alignment(vertical='top')
        ws.cell(row=r_idx, column=5, value=pcs).alignment = Alignment(vertical='top')
        ws.cell(row=r_idx, column=6, value=status).alignment = Alignment(vertical='top')

        c7 = ws.cell(row=r_idx, column=7)
        if is_url(service):
            c7.value = simplify_service_for_display(service)
            c7.hyperlink = service
            c7.font = Font(color='0000FF', underline='single')
        else:
            c7.value = service
        c7.alignment = Alignment(vertical='top')

        for c in range(1, 8):
            ws.cell(row=r_idx, column=c).border = THIN_BORDER

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'A2'

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

def create_manifest_zip_bundle(files_dict: Dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for fname, fbytes in files_dict.items():
            if fbytes:
                z.writestr(fname, fbytes)
    out.seek(0)
    return out.getvalue()

def extract_sheet_matrix(excel_bytes: bytes) -> Dict[str, List[List[Any]]]:
    """Read back an Excel workbook in memory and return each sheet as a 2D array of cell values."""
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(['' if v is None else v for v in row])
        while rows and all(c == '' for c in rows[-1]):
            rows.pop()
        sheets[name] = rows
    return sheets

def extract_chamber_preview(docx_bytes: bytes) -> Dict[str, Any]:
    """Extract text content and tables from chamber certificate docx for preview."""
    if not docx_bytes:
        return {"paragraphs": [], "tables": []}
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip() != '']
    tables = []
    for table in doc.tables:
        tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
    return {"paragraphs": paragraphs, "tables": tables}

